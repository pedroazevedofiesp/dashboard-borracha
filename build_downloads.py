#!/usr/bin/env python3
"""Gera as planilhas Excel formatadas de data/downloads/, uma por seção do
painel (uma para cada CNAE de config_painel.json onde a seção varia por
segmento). Lê direto dos CSVs de 3_Dados_Tratados_CSV/ — roda depois de
build_data.py, mas não depende do data.json."""

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_data import (CNAES, CONFIG, ROTULO_ARQUIVO, SETOR, SLUG, UF_REGIAO, build_fator_ipca, build_macro,
                        tarifa_industrial_regiao_ano)

ROOT = Path(__file__).resolve().parent
SRC_DIR = ROOT / '3_Dados_Tratados_CSV'
OUT_DIR = ROOT / 'data' / 'downloads'

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='000000')
BODY_FONT = Font(name=FONT_NAME, color='000000')
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color='595959')
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=13, color='000000')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NO_SEP_HINTS = ('ano', 'mes', 'mês', 'cnae', 'codigo', 'código', 'code', 'reporter_code', 'partner_code')


def _read(name, **kwargs):
    defaults = dict(sep=';', decimal=',', encoding='utf-8-sig')
    defaults.update(kwargs)
    return pd.read_csv(SRC_DIR / name, **defaults)


def _sheet_name(name):
    name = re.sub(r'[:\\/?*\[\]]', '-', str(name))
    return name[:31]


def _is_plain_int_col(header):
    h = str(header).lower()
    return any(hint in h for hint in NO_SEP_HINTS)



# Acima disso, estilo por célula (borda + fonte em cada uma) fica lento
# demais (Comtrade e energia industrial passam de 100 mil linhas) — nesse
# caso usa escrita em bloco (ws.append) e só formata número, sem borda/fonte
# por célula, pra não travar o build.
LARGE_SHEET_THRESHOLD = 3000


def _normalize_value(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _write_df(ws, df, title=None, note=None):
    r = 1
    if title:
        ws.cell(row=r, column=1, value=title).font = TITLE_FONT
        r += 2
    headers = list(df.columns)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=c, value=str(h))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    header_row = r
    r += 1
    plain_int_cols = {h for h in headers if _is_plain_int_col(h)}
    n_rows = len(df)
    style_cells = n_rows <= LARGE_SHEET_THRESHOLD

    if style_cells:
        for _, row in df.iterrows():
            for c, h in enumerate(headers, start=1):
                v = _normalize_value(row[h])
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = BODY_FONT
                cell.border = BORDER
                if isinstance(v, (int, float)):
                    cell.number_format = '0' if h in plain_int_cols else '#,##0.##'
            r += 1
    else:
        # values.tolist() + ws.append: ordens de magnitude mais rápido que
        # ws.cell() célula a célula pra planilhas de centenas de milhares
        # de linhas.
        for row_vals in df[headers].values.tolist():
            ws.append([_normalize_value(v) for v in row_vals])
        num_formats = ['0' if h in plain_int_cols else '#,##0.##' for h in headers]
        for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + n_rows):
            for cell, fmt in zip(row, num_formats):
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
        r = header_row + 1 + n_rows

    for c, h in enumerate(headers, start=1):
        sample = df[h].head(300)
        max_len = max((len(str(x)) for x in sample if pd.notna(x)), default=10)
        width = max(10, min(42, len(str(h)) + 2, max_len + 2))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = f'A{header_row + 1}'
    if note:
        nr = r + 1
        ws.cell(row=nr, column=1, value=note).font = NOTE_FONT
        ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=max(1, len(headers)))
    return r


def _new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _add_sheet(wb, sheet_name, df, title=None, note=None):
    base = _sheet_name(sheet_name)
    name = base
    i = 2
    while name in wb.sheetnames:
        name = _sheet_name(f'{base[:28]}_{i}')
        i += 1
    ws = wb.create_sheet(name)
    _write_df(ws, df, title=title, note=note)
    return ws


def _save(wb, filename):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename
    wb.save(path)
    print(f'  data/downloads/{filename}')
    return path


# ---------------------------------------------------------------------------
# 01 — Produção física (índice da divisão; igual para todos os segmentos,
# por isso um arquivo só)
# ---------------------------------------------------------------------------
def build_producao_download():
    wb = _new_workbook()
    _add_sheet(wb, 'Indice mensal', _read(f"Producao_Fisica_Mensal_{SETOR['producao_rotulo_arquivo']}.csv"),
               title=f"Produção física — {SETOR['producao_titulo']} (índice mensal, base 100)",
               note=f"Fonte: IBGE/PIM-PF. {SETOR['producao_descricao']}")
    _save(wb, 'Producao_Fisica.xlsx')


# ---------------------------------------------------------------------------
# 02 — Financeiro (PIA, não varia por setor)
# ---------------------------------------------------------------------------
def build_financeiro_download():
    grupo = SETOR['grupo_cnae']
    codigos = '_'.join(c['codigo'] for c in CNAES)
    rotulos = ' x '.join(f"{c['rotulo']} ({c['codigo']})" for c in CNAES)
    wb = _new_workbook()
    _add_sheet(wb, f'Grupo {grupo}', _read(f"Dados_Financeiros_{ROTULO_ARQUIVO}_{grupo.replace('.', '_')}.csv"),
               title=f"Financeiro — {SETOR['grupo_descricao']} (CNAE {grupo})",
               note=f"Grupo CNAE completo, sem abertura por classe. {SETOR.get('grupo_nota', '')} "
                    f"Fonte: IBGE/PIA-Empresa.")
    _add_sheet(wb, 'Por CNAE (2024)', _read(f'Dados_Financeiros_CNAE_{codigos}_2024_SerieNova.csv'),
               title=f'Financeiro — {rotulos}',
               note='Só 2024 — primeira vez com essa abertura por classe. Fonte: IBGE/PIA-Empresa.')
    _save(wb, 'Financeiro.xlsx')


# ---------------------------------------------------------------------------
# 03 — Emprego formal / RAIS (por setor)
# ---------------------------------------------------------------------------
def emprego_files(cfg):
    cnae = cfg['codigo']
    rotulo = cfg.get('rotulo_arquivo', ROTULO_ARQUIVO)
    return [
        (f'Empregos_RAIS_UF_{rotulo}_{cnae}.csv', 'RAIS por UF'),
        (f'Empregos_RAIS_Tamanho_{rotulo}_{cnae}.csv', 'RAIS por tamanho'),
        (f'rais_vinc_{SLUG}_escolaridade_{cnae}.csv', 'RAIS escolaridade'),
        (f'rais_vinc_{SLUG}_massa_{cnae}.csv', 'RAIS massa salarial'),
        (f'rais_vinc_{SLUG}_ocupacao_{cnae}.csv', 'RAIS ocupacao detalhada'),
        (f'rais_vinc_{SLUG}_ocupacao_agrupada_{cnae}.csv', 'RAIS ocupacao agrupada'),
        (f'rais_vinc_{SLUG}_tempo_emprego_{cnae}.csv', 'RAIS tempo de emprego'),
    ]


def build_emprego_download(cfg):
    wb = _new_workbook()
    for fname, sheet in emprego_files(cfg):
        _add_sheet(wb, sheet, _read(fname), title=f"Emprego formal (RAIS) — {cfg['rotulo']}",
                   note='Fonte: RAIS.')
    _save(wb, f"Emprego_{cfg['codigo']}.xlsx")


# ---------------------------------------------------------------------------
# 04 — CAGED (arquivos únicos com os dois setores, filtrados por cnae_subclasse)
# ---------------------------------------------------------------------------
def build_caged_download(cfg):
    cnae, label = cfg['codigo'], cfg['rotulo']
    wb = _new_workbook()
    for fname, sheet in [
        (f'caged_{SLUG}_saldo_mensal_uf_cnae.csv', 'CAGED saldo mensal'),
        (f'caged_{SLUG}_salario_mensal_uf_cnae.csv', 'CAGED salario mensal'),
        (f'caged_{SLUG}_tipo_movimentacao_mensal_cnae.csv', 'CAGED tipo movimentacao'),
    ]:
        df = _read(fname)
        df = df[df['cnae_subclasse'].astype(str).str.strip() == cnae].reset_index(drop=True)
        _add_sheet(wb, sheet, df, title=f'CAGED — {label}', note='Fonte: CAGED/Novo CAGED.')
    _save(wb, f'CAGED_{cnae}.xlsx')


# ---------------------------------------------------------------------------
# 05 — Comércio exterior (por setor: cesta de NCM, país, Comtrade)
# ---------------------------------------------------------------------------
def build_comex_download(cfg):
    cnae, label = cfg['codigo'], cfg['rotulo']
    rotulo = cfg.get('rotulo_arquivo', ROTULO_ARQUIVO)
    cesta = f"Cesta de NCMs do segmento: {cfg['cesta_ncm']}"
    wb = _new_workbook()
    _add_sheet(wb, 'Balanca por NCM', _read(f'Comex_{cnae}_NCM.csv'),
               title=f'Comércio exterior — Balança comercial por NCM — {label}',
               note=f'Fonte: Comex Stat/MDIC. {cesta}')
    _add_sheet(wb, 'UF por NCM', _read(f'Comex_{cnae}_NCM_UF.csv'),
               title=f'Comércio exterior — UF de origem, por NCM — {label}',
               note='Mesma cesta de NCMs da aba "Balanca por NCM" — usada no ranking por UF do painel.')
    _add_sheet(wb, 'Pais parceiro', _read(f'Comex_Exportacao_Importacao_{cnae}.csv'),
               title=f'Comércio exterior — País parceiro e UF — {label}',
               note='Mesma cesta de NCMs, aberta por país — usada no ranking por país do painel.')
    _add_sheet(wb, 'Comtrade mundial', _read(f'Comtrade_Global_{rotulo}_{cnae}.csv'),
               title=f'Comércio mundial (Comtrade, proxy por HS) — {label}',
               note='Contexto global — classificação por código HS como proxy do CNAE, não é o número oficial de comércio exterior do Brasil (ver Comex/MDIC pra isso).')
    _save(wb, f'Comercio_Exterior_{cnae}.xlsx')


# ---------------------------------------------------------------------------
# 06 — BNDES (por setor)
# ---------------------------------------------------------------------------
def build_bndes_download(cfg):
    cnae = cfg['codigo']
    df = _read(f'BNDES_Desembolsos_{cnae}.csv')
    wb = _new_workbook()
    _add_sheet(wb, 'Desembolsos', df,
               title=f"BNDES — Desembolsos — {cfg['rotulo']}",
               note=f"Fonte: BNDES, {int(df['Ano'].min())}-{int(df['Ano'].max())}. Valores em R$ correntes.")
    _save(wb, f'BNDES_{cnae}.xlsx')


# ---------------------------------------------------------------------------
# 07 — DECOM (bloco opcional, só se config_painel.json apontar o arquivo)
# ---------------------------------------------------------------------------
def build_decom_download(arquivo):
    wb = _new_workbook()
    _add_sheet(wb, 'Processos', _read(arquivo),
               title='DECOM — Processos de defesa comercial',
               note='Fonte: DECOM/GECEX.')
    _save(wb, 'DECOM.xlsx')


# ---------------------------------------------------------------------------
# Energia Industrial (view própria, comum a todos os painéis setoriais)
# ---------------------------------------------------------------------------
def build_energia_download():
    wb = _new_workbook()

    # Custo = tarifa média REAL faturada da classe Industrial por região
    # (MME/ANEEL), não o fator-de-carga-assumido do CSV bruto (checado contra
    # tarifa publicada e descartado — dava menos da metade do valor real).
    # Gasto/custo real (deflacionado por IPCA) na mesma lógica do build_data.py.
    df_energia = _read('energia_industria_transformacao_estados_brasil_2012-2026.csv', decimal='.')
    df_energia = df_energia.drop(columns=['FATOR_CARGA_ASSUMIDO'])  # não alimenta mais o custo, ver nota abaixo
    macro = build_macro()
    ipca_by_year = {}
    for r in macro:
        if r['ipca'] is not None:
            ipca_by_year.setdefault(r['ano'], []).append(r['ipca'])
    ipca_avg = {ano: sum(v) / len(v) for ano, v in ipca_by_year.items()}

    df_energia['CONSUMO_MWH'] = pd.to_numeric(df_energia['CONSUMO_MWH'], errors='coerce')
    df_energia['REGIAO'] = df_energia['UF'].map(UF_REGIAO).fillna('Brasil')
    df_energia['CUSTO_TOTAL_RS_MWH'] = df_energia.apply(
        lambda r: tarifa_industrial_regiao_ano(r['REGIAO'], int(r['ANO']), ipca_avg), axis=1)
    df_energia['GASTO_ESTIMADO_RS'] = df_energia['CONSUMO_MWH'] * df_energia['CUSTO_TOTAL_RS_MWH']

    meses_por_ano = df_energia.groupby('ANO')['MES'].nunique()
    ano_max = int(df_energia['ANO'].max())
    ano_base_energia = ano_max if meses_por_ano.get(ano_max, 0) >= 12 else ano_max - 1
    fator_ipca_energia = build_fator_ipca(macro, ano_base_energia)
    df_energia['FATOR_IPCA'] = df_energia['ANO'].map(fator_ipca_energia)
    df_energia[f'CUSTO_REAL_RS_MWH_{ano_base_energia}'] = df_energia['CUSTO_TOTAL_RS_MWH'] * df_energia['FATOR_IPCA']
    df_energia[f'GASTO_ESTIMADO_REAL_RS_{ano_base_energia}'] = df_energia['GASTO_ESTIMADO_RS'] * df_energia['FATOR_IPCA']
    df_energia = df_energia.drop(columns=['FATOR_IPCA'])

    _add_sheet(wb, 'Estados x divisoes CNAE', df_energia,
               title='Energia industrial — Consumo/custo por UF e divisão CNAE (2012-2026)',
               note=f'Consumo (MWh): fonte própria por UF/divisão CNAE. Custo (R$/MWh): tarifa média REAL faturada '
                    f'da classe Industrial por região — MME/ANEEL, "Informativo Tarifário/Gestão do Setor Elétrico", '
                    f'Tabela "Tarifa Média Faturada por Classe de Consumo e Região" (gov.br/mme). Não varia por '
                    f'estado dentro da região nem por divisão CNAE. Anos sem edição publicada são interpolados/'
                    f'extrapolados por IPCA (ver coluna REGIAO). Gasto/custo real em R$ de {ano_base_energia} '
                    f'(deflacionado por IPCA).')
    _add_sheet(wb, 'SP detalhado', _read('energia_industria_transformacao_sp_brasil_2012-2026.csv', decimal='.'),
               title='Energia industrial — São Paulo (detalhado)')
    _save(wb, 'Energia_Industrial.xlsx')


def build_all_downloads():
    print('Gerando planilhas de download (data/downloads/)...')
    build_producao_download()
    build_financeiro_download()
    for cfg in CNAES:
        build_emprego_download(cfg)
        build_caged_download(cfg)
        build_comex_download(cfg)
        build_bndes_download(cfg)
    decom = CONFIG.get('blocos_opcionais', {}).get('decom')
    if decom:
        build_decom_download(decom)
    build_energia_download()


if __name__ == '__main__':
    build_all_downloads()
